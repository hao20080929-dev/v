import os
import re
import sys
import time
from datetime import date
import requests
import urllib3
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from google import genai

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
time.sleep(2)


def build_prompt(property_description: str) -> str:
    return (
        "你是專業在地房仲，講話直白，避免「擁抱幸福、溫馨家園、祝您成功、荷包、夢幻、精準、安居樂業」等情緒詞。"
        "請使用「利潤、稀缺、抗跌、起家厝」等強勢詞彙。"
        "請使用台灣慣用語，必須出現坪數、陽台、坡平車位等詞彙，且不得出現中國大陸術語。"
        "請判斷行政區（雙北/中南部/重劃區）並適配價值主張（捷運/學區/台積電/低總價）。"
        "若輸入包含總價與坪數，請自動計算單價；若缺資料，請明確標示「未提供」。"
        "請找出物件周邊 3 個具體生活地標，僅保留「捷運、明星學區、大型商場、重大建設」，避免「便利商店、加油站」等通用設施。"
        "請濃縮輸出，避免重複，整體長度以 3-4 頁內為目標。"
        "輸出包含以下內容（每區塊控制精簡）："
        "\n1. 數據深度偵察（Data Intelligence）："
        "\n- 行情對比：分析單價並與行政區平均行情比對，給出「比區域行情低 X%」或「預售鎖利空間 X 萬」。"
        "\n- 持有成本分析：無管理費或公設比低強調「養房無負擔」；預售屋強調「工程期付款節奏」。"
        "\n2. 房仲專屬三合一輸出（不得省略）："
        "\n- 【591 專業版】：數據極致精準、條列式優點、專業地標連結（明星學區、捷運出口、台積電廠區預定地）。"
        "\n- 【FB 社團吸粉版】：第一句強力 Hook，在地化口語＋大量 Emoji，強調鄰里素質與生活機能。"
        "\n- 【LINE/限動秒殺版】：不超過 100 字，強調屋主急售、已有多組預約、晚了就沒了。"
        "\n3. 心理誘發技術："
        "\n- 自住客：營造幸福感與安全感（孩子走路就上學、太太買菜不累）。"
        "\n- 投資客：強調稀缺性與增值槓桿（收租投報高、區域開發紅利）。"
        "\n4. 法律與數據防線："
        "\n- 只保留單行警示字，避免長篇免責。"
        "\n5. 📸 專業配圖建議：提供 3 種最吸睛的照片題材。"
        f"\n待處理資訊：{property_description}"
    )


def resolve_input_text(user_input: str) -> tuple[str, str]:
    if "591.com.tw" not in user_input:
        return user_input, ""
    print("正在解析網址...")
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/125.0 Safari/537.36",
    }
    response = requests.get(user_input, headers=headers, timeout=15, verify=False)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else ""
    text = soup.get_text(separator=" ", strip=True)
    return text[:2000], title


def sanitize_filename(value: str) -> str:
    cleaned = re.sub(r"[\\/:*?\"<>|]", "_", value)
    cleaned = re.sub(r"\s+", "_", cleaned).strip("_")
    return cleaned or "物件"


def split_inputs(raw_input: str) -> list[str]:
    parts = re.split(r"[,\n]+", raw_input)
    return [item.strip() for item in parts if item.strip()]


def extract_sections(content: str) -> dict[str, str]:
    sections = {
        "【591 專業版】": "",
        "【FB 社團吸粉版】": "",
        "【LINE/限動秒殺版】": "",
    }
    current_key = None
    for line in content.splitlines():
        line_strip = line.strip()
        if line_strip in sections:
            current_key = line_strip
            continue
        if current_key:
            sections[current_key] += line + "\n"
    return sections


def filter_generic_landmarks(lines: list[str]) -> list[str]:
    blocked = ["便利商店", "加油站"]
    filtered: list[str] = []
    for line in lines:
        if any(keyword in line for keyword in blocked):
            continue
        filtered.append(line)
    return filtered


def limit_lines(lines: list[str], max_lines: int) -> list[str]:
    return lines[:max_lines]


def extract_intel_section(content: str) -> str:
    start_marker = "1. 數據深度偵察"
    end_marker = "2. 房仲專屬三合一輸出"
    start_index = content.find(start_marker)
    if start_index == -1:
        return content
    end_index = content.find(end_marker, start_index)
    if end_index == -1:
        return content[start_index:]
    return content[start_index:end_index]


def clean_markdown(text: str) -> str:
    return re.sub(r"\*+", "", text)


def add_field(paragraph, instruction: str) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    paragraph._p.append(field)


def add_paragraph_with_highlight(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pattern = re.compile(r"(總價[^，。\n]*|單價[^，。\n]*|完工日期[^，。\n]*)")
    remaining = text
    last_index = 0
    for match in pattern.finditer(text):
        start, end = match.span()
        if start > last_index:
            paragraph.add_run(text[last_index:start])
        highlighted = paragraph.add_run(text[start:end])
        highlighted.font.highlight_color = WD_COLOR_INDEX.YELLOW
        last_index = end
    if last_index < len(text):
        paragraph.add_run(text[last_index:])


def apply_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft JhengHei"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    heading1 = document.styles["Heading 1"]
    heading1.font.name = "Microsoft JhengHei"
    heading1.font.size = Pt(18)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor(0, 51, 102)
    heading1._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    heading1._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    heading2 = document.styles["Heading 2"]
    heading2.font.name = "Microsoft JhengHei"
    heading2.font.size = Pt(14)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor(0, 0, 0)
    heading2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    heading2._element.rPr.rFonts.set(qn("w:ascii"), "Arial")


def set_header_footer(document: Document, property_name: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    header = section.header
    header_table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    header_table.autofit = True
    left_cell = header_table.cell(0, 0)
    right_cell = header_table.cell(0, 1)
    left_cell.text = f"[{property_name}]"
    right_cell.text = date.today().strftime("%Y/%m/%d")
    right_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_paragraph.add_run("第 ")
    add_field(footer_paragraph, "PAGE")
    footer_paragraph.add_run(" 頁 / 共 ")
    add_field(footer_paragraph, "NUMPAGES")
    footer_paragraph.add_run(" 頁｜⚠️ 數據需人工確認")


def extract_region_and_name(title: str, fallback: str) -> tuple[str, str]:
    region_candidates = [
        "台北市", "新北市", "桃園市", "台中市", "台南市", "高雄市", "基隆市", "新竹市", "嘉義市",
        "台北", "新北", "桃園", "台中", "台南", "高雄", "基隆", "新竹", "嘉義",
        "宜蘭", "花蓮", "台東", "雲林", "彰化", "南投", "屏東", "苗栗",
    ]
    region = "全台"
    for candidate in region_candidates:
        if candidate in title:
            region = candidate
            break
    name_source = title or fallback
    name_source = re.sub(r"591|租|售|買|房屋|物件|出售", "", name_source)
    name_source = re.sub(r"[\[\]\(\)【】]", "", name_source)
    name = sanitize_filename(name_source)[:12] or "物件"
    return region, name


def generate_listing(client: genai.Client, property_description: str) -> str:
    prompt = build_prompt(property_description)
    model_sequence = [
        "models/gemini-2.5-flash",
        "models/gemini-2.0-flash",
        "models/gemini-1.5-flash",
        "models/gemma-3-12b",
    ]
    last_error: Exception | None = None
    for model_name in model_sequence:
        try:
            response = client.models.generate_content(
                model=model_name,
                contents=prompt,
                config={"temperature": 0.7},
            )
            content = getattr(response, "text", None)
            if content:
                return content.strip()
            last_error = ValueError("Gemini response content is empty")
        except Exception as error:
            last_error = error
            continue
    if last_error:
        raise last_error
    raise ValueError("Gemini response content is empty")


def save_docx(title: str, content: str) -> str:
    output_dir = os.path.join(os.getcwd(), "Outputs")
    os.makedirs(output_dir, exist_ok=True)
    cleaned_content = clean_markdown(content)
    intel_section = clean_markdown(extract_intel_section(cleaned_content))
    sections = extract_sections(cleaned_content)
    document = Document()
    apply_styles(document)
    region, property_name = extract_region_and_name(title, intel_section[:12])
    set_header_footer(document, property_name)

    document.add_heading("行情深度偵察", level=1)
    document.add_paragraph("#行情對比 #抗跌 #稀缺")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "本案單價"
    table.cell(0, 1).text = "區域行情"
    table.cell(1, 0).text = "價差判讀"
    table.cell(1, 1).text = "待 AI 解析"
    intel_lines = limit_lines(filter_generic_landmarks(intel_section.splitlines()), 12)
    for line in intel_lines:
        line_strip = line.strip()
        if line_strip:
            add_paragraph_with_highlight(document, line_strip)
    document.add_page_break()

    document.add_heading("591 專業版", level=1)
    document.add_paragraph("#成交戰術 #數據精準 #專業建議")
    professional_lines = limit_lines(filter_generic_landmarks((sections["【591 專業版】"].strip() or cleaned_content).splitlines()), 16)
    for line in professional_lines:
        line_strip = line.strip()
        if line_strip:
            add_paragraph_with_highlight(document, line_strip)
    document.add_page_break()

    document.add_heading("FB 社團吸粉版", level=1)
    document.add_paragraph("#在地社群 #吸粉曝光 #熱區生活")
    fb_lines = limit_lines(filter_generic_landmarks(sections["【FB 社團吸粉版】"].strip().splitlines()), 10)
    for line in fb_lines:
        line_strip = line.strip()
        if line_strip:
            add_paragraph_with_highlight(document, line_strip)
    document.add_page_break()

    document.add_heading("LINE/限動秒殺版", level=1)
    document.add_paragraph("#VIP急售 #限量釋出 #稀缺搶手")
    line_lines = limit_lines(filter_generic_landmarks(sections["【LINE/限動秒殺版】"].strip().splitlines()), 6)
    for line in line_lines:
        line_strip = line.strip()
        if line_strip:
            add_paragraph_with_highlight(document, line_strip)
    filename = f"【戰術報告】{region}_{property_name}_{date.today().strftime('%m%d')}.docx"
    file_path = os.path.join(output_dir, filename)
    document.save(file_path)
    return file_path


def main() -> None:
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise SystemExit("請設定 GEMINI_API_KEY")
    client = genai.Client(api_key=api_key)
    user_input = input("請直接貼上 591 網址或房屋資訊：")
    inputs = split_inputs(user_input)
    if not inputs:
        raise SystemExit("請提供房地產物件描述文字")
    for item in inputs:
        description, title = resolve_input_text(item)
        if not description:
            continue
        output = generate_listing(client, description)
        base_title = title or (description[:12] if description else "物件")
        save_docx(base_title, output)
        print("✅ 超級專家報告已生成在 Outputs 資料夾")


if __name__ == "__main__":
    main()
